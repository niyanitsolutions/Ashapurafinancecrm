"""Bounded in-memory parsing. Files are never persisted or evaluated."""

import csv
import io
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import PurePath
from typing import Any
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse  # type: ignore[import-untyped]
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.cell import (  # type: ignore[import-untyped]
    column_index_from_string,
    coordinate_from_string,
)
from pydantic import BaseModel

from app.core.exceptions import ValidationError
from app.features.insurance_management.schemas import CreateManualInsuranceCaseRequest
from app.features.leads.schemas import CreateLeadRequest

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
MODELS: dict[str, type[BaseModel]] = {
    "leads": CreateLeadRequest,
    "insurance": CreateManualInsuranceCaseRequest,
}
EXCLUDED = {
    "latitude",
    "longitude",
    "product_form_data",
    "stage",
    "reason",
    "re_eligibility",
    "re_eligible_date",
}
LABELS = {
    "source_id": "Lead Source",
    "product_id": "Product",
    "insurance_category_id": "Insurance Category",
    "preferred_amount": "Preferred Loan Amount",
    "comment": "Comments",
    "assigned_to": "Assign To",
    "nominee_dob": "Nominee Date Of Birth",
}


def fields(kind: str) -> dict[str, str]:
    return {
        key: LABELS.get(key, key.replace("_", " ").title())
        for key in MODELS[kind].model_fields
        if key not in EXCLUDED
    }


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_file(
    content: bytes, filename: str, content_type: str | None, kind: str
) -> list[tuple[int, dict[str, str]]]:
    extension = PurePath(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise ValidationError(
            "Choose a CSV or XLSX file. XLS and macro-enabled files are not supported."
        )
    if not content or len(content) > MAX_BYTES:
        raise ValidationError("Choose a non-empty file no larger than 5 MB.")
    allowed_types = {None, "", "application/octet-stream"}
    allowed_types |= (
        {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"}
        if extension == ".csv"
        else {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }
    )
    if content_type not in allowed_types:
        raise ValidationError("The file type does not match CSV or XLSX.")
    workbook = None
    rows: Iterator[Any]
    try:
        if extension == ".csv":
            decoded = content.decode("utf-8-sig")
            if "\x00" in decoded:
                raise ValueError("binary CSV")
            rows = csv.reader(io.StringIO(decoded, newline=""), strict=True)
        else:
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 2000
                    or sum(entry.file_size for entry in entries) > 25 * 1024 * 1024
                ):
                    raise ValueError("oversized workbook")
                if any(
                    entry.flag_bits & 1
                    or "vbaproject" in entry.filename.lower()
                    or "externallinks/" in entry.filename.lower()
                    for entry in entries
                ):
                    raise ValueError("unsupported workbook")
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            sheet = workbook.worksheets[0]
            # Validate actual coordinates before openpyxl can expand sparse cells into
            # huge rows. Dimensions supplied by a workbook producer are not trusted.
            with (
                ZipFile(io.BytesIO(content)) as archive,
                archive.open(sheet._worksheet_path) as xml,
            ):
                for _event, element in iterparse(xml, events=("end",), forbid_dtd=True):
                    if (
                        element.tag.endswith("}row")
                        and int(element.attrib.get("r", "0")) > MAX_ROWS + 1
                    ):
                        raise ValidationError(
                            "The first worksheet exceeds the supported 500 records."
                        )
                    if element.tag.endswith("}c"):
                        column, number = coordinate_from_string(element.attrib.get("r", ""))
                        if number > MAX_ROWS + 1 or column_index_from_string(column) > len(
                            fields(kind)
                        ):
                            raise ValidationError(
                                "The first worksheet exceeds the supported 500 records or sample columns."
                            )
                    element.clear()
            sheet.reset_dimensions()  # Do not trust producer-provided dimensions.
            rows = (
                [cell.value if cell.data_type != "f" else "=FORMULA_NOT_SUPPORTED" for cell in row]
                for row in sheet.iter_rows()
            )
        header = next(rows, None)
        if header is None or len(header) > len(fields(kind)):
            raise ValueError("header")
        mapping = {label.casefold(): key for key, label in fields(kind).items()}
        mapping.update({key.casefold(): key for key in fields(kind)})
        keys = [mapping.get(cell_text(value).strip().casefold()) for value in header]
        if not keys or None in keys or len(set(keys)) != len(keys):
            raise ValidationError(
                "Use the sample column names. Unknown, blank, or repeated columns were found."
            )
        keys = [key for key in keys if key is not None]
        required = {
            key
            for key, field in MODELS[kind].model_fields.items()
            if field.is_required() and key not in EXCLUDED
        }
        if not required.issubset(keys):
            raise ValidationError(
                "Missing required columns: "
                + ", ".join(fields(kind)[key] for key in sorted(required - set(keys)))
            )
        result = []
        for number, row in enumerate(rows, 2):
            if number > MAX_ROWS + 1:
                raise ValidationError(
                    "A file can contain at most 500 rows. Split it into smaller files."
                )
            if len(row) > len(keys):
                raise ValidationError(f"Row {number} contains extra columns.")
            values = [cell_text(value) for value in row]
            if any(len(value) > 10000 for value in values):
                raise ValidationError(
                    f"Row {number} contains a cell longer than 10,000 characters."
                )
            # Blank physical rows remain visible as invalid rows; never silently drop data.
            result.append(
                (
                    number,
                    {
                        key: value
                        for key, value in zip(keys, values, strict=False)
                        if key is not None
                    },
                )
            )
        if not result:
            raise ValidationError("The file contains no records.")
        return result
    except ValidationError:
        raise
    except (
        ValueError,
        TypeError,
        KeyError,
        IndexError,
        UnicodeError,
        csv.Error,
        BadZipFile,
        OSError,
    ):
        raise ValidationError(
            "Unable to process the file. Please check the file format and try again."
        ) from None
    except Exception:  # noqa: BLE001 -- third-party malformed XML must produce a safe file error
        # XML parser errors must not expose workbook content or implementation details.
        raise ValidationError(
            "Unable to process the file. Please check the file format and try again."
        ) from None
    finally:
        if workbook is not None:
            workbook.close()


def sample_file(kind: str, lookups: dict[str, list[dict[str, str]]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    columns = fields(kind)
    sheet.append(list(columns.values()))
    sheet.freeze_panes = "A2"
    example = {"full_name": "Ramesh Kumar", "mobile": "9876543210", "email": "ramesh@example.com"}
    if kind == "leads":
        example["product_category"] = "loan"
        for field, catalog in (("source_id", "sources"), ("product_id", "loan_products")):
            example[field] = (
                lookups[catalog][0]["name"] if lookups[catalog] else "Select from Lookups"
            )
    else:
        product = next(iter(lookups["insurance_products"]), None)
        example["product_id"] = product["name"] if product else "Select from Lookups"
        example["insurance_category_id"] = next(
            (
                item["name"]
                for item in lookups["categories"]
                if product and item["id"] == product.get("category_id")
            ),
            "Select from Lookups",
        )
        example["gender"] = "male"
    sheet.append([example.get(key, "") for key in columns])
    instructions = workbook.create_sheet("Instructions")
    for text in (
        "Replace the example row with your own records. Only the first sheet is imported. Maximum 500 rows / 5 MB.",
        "Source, Product and Insurance Category accept an exact name or ID from Lookups. Ambiguous names require an ID.",
        "Dates: YYYY-MM-DD or DD-MM-YYYY. Mobile: 10 digits, starting with 6, 7, 8 or 9. Store mobile numbers as text.",
        "Gender: male, female or other. Product Category: loan or insurance. Empty optional cells are omitted.",
        "Cells beginning with =, +, - or @ are rejected as spreadsheet formula prefixes. Use plain text values.",
        "Assign To: an eligible employee ID or 'self'. Leave empty for Fresh Leads. Assignment requires your existing Assign permission.",
        "General Leads reject active duplicate mobiles. Insurance permits additional applications for an existing customer, as in manual creation.",
        "Preview every row before confirming. Reusing the same file within 24 hours returns the same batch and cannot import it twice.",
        "Required columns: "
        + ", ".join(
            columns[key]
            for key, field in MODELS[kind].model_fields.items()
            if key in columns and field.is_required()
        ),
    ):
        instructions.append([text])
    lookup_sheet = workbook.create_sheet("Lookups")
    lookup_sheet.append(["List", "Name", "ID", "Category ID"])
    for catalog, items in lookups.items():
        for item in items:
            lookup_sheet.append([catalog, item["name"], item["id"], item.get("category_id", "")])
    # Even a master-data name beginning '=' must be written as text, never a formula.
    for page in workbook:
        for row in page:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
