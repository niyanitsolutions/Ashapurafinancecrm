"""Local mocked-database coverage of preview, confirmation and normal creation paths."""

import base64
import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from test_insurance_employee_visibility import _employee_with_permissions
from test_insurance_manual_lead import _payload
from test_insurance_policy_leads import _insurance_product_with_schema
from test_leads import _create_employee, _grant_leads_actions, _lead_master_data, _lead_payload
from test_workflow import _login, _seed_workflow_definitions

from app.features.leads.service import LeadService


async def test_authorized_employees_import_and_cannot_use_another_preview(
    client, mock_db, owner_headers, master_data
):
    master = await _lead_master_data(mock_db)
    employee = await _create_employee(client, owner_headers, master_data)
    await _grant_leads_actions(
        client, owner_headers, employee["id"], ["view", "create"], "Importer"
    )
    headers = await _login(client, "9511111111", "InitialPass1!")
    batch = await preview(client, headers, [_lead_payload(master)])
    response = await client.post(
        "/api/v1/bulk-import/leads/confirm",
        json={"batch_id": batch["batch_id"]},
        headers=owner_headers,
    )
    assert response.status_code == 404
    assert (await confirm(client, headers, batch))["imported"] == 1
    assignment = await preview(
        client, headers, [_lead_payload(master, mobile="9611111122", assigned_to="self")]
    )
    assert assignment["invalid"] == 1
    assert "Assign permission" in assignment["rows"][0]["error"]
    await _seed_workflow_definitions(mock_db)
    product = await _insurance_product_with_schema(
        mock_db, product_name="Employee Plan", docs=["PAN"]
    )
    _employee, insurance_headers = await _employee_with_permissions(
        client,
        owner_headers,
        master_data,
        mobile="9811000095",
        email="import@example.com",
        actions=["view", "edit"],
    )
    row = _payload(product)
    row.pop("stage")
    insurance_batch = await preview(client, insurance_headers, [row], "insurance")
    assert (await confirm(client, insurance_headers, insurance_batch, "insurance"))["imported"] == 1


async def test_confirm_rechecks_existing_duplicates(client, mock_db, owner_headers):
    master = await _lead_master_data(mock_db)
    payload = _lead_payload(master)
    batch = await preview(client, owner_headers, [payload])
    assert (
        await client.post("/api/v1/leads", json=payload, headers=owner_headers)
    ).status_code == 200
    result = await confirm(client, owner_headers, batch)
    assert (result["duplicate"], result["imported"]) == (1, 0)


async def test_sparse_workbook_is_rejected(client, owner_headers):
    workbook = Workbook()
    workbook.active["A1000000"] = "far away"
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = await client.post(
        "/api/v1/bulk-import/leads/preview",
        files={
            "file": (
                "sparse.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=owner_headers,
    )
    assert response.status_code == 422


async def test_large_batch_resumes_only_unprocessed_rows(client, mock_db, owner_headers):
    master = await _lead_master_data(mock_db)
    rows = [_lead_payload(master, mobile=f"96{number:08d}") for number in range(30)]
    batch = await preview(client, owner_headers, rows)
    first = await confirm(client, owner_headers, batch)
    assert (first["imported"], first["valid"], first["state"]) == (25, 5, "preview")
    result = await confirm(client, owner_headers, batch)
    assert (result["imported"], result["valid"], result["state"]) == (30, 0, "completed")
    assert (await confirm(client, owner_headers, batch))["imported"] == 30
    assert await mock_db.leads.count_documents({}) == 30


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
async def test_csv_formula_prefixes_are_reported_without_execution(
    client, mock_db, owner_headers, prefix
):
    master = await _lead_master_data(mock_db)
    batch = await preview(
        client, owner_headers, [_lead_payload(master, remarks=f"{prefix}SUM(1,2)")]
    )
    assert batch["invalid"] == 1
    assert "formulas" in batch["rows"][0]["error"]
    assert await mock_db.leads.count_documents({}) == 0


def upload(rows, extension="csv"):
    keys = list(dict.fromkeys(key for row in rows for key in row))
    if extension == "csv":
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)
        return ("leads.csv", buffer.getvalue().encode(), "text/csv")
    workbook = Workbook()
    workbook.active.append(keys)
    for row in rows:
        workbook.active.append([row.get(key) for key in keys])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return (
        "leads.xlsx",
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


async def preview(client, headers, rows, kind="leads", extension="csv"):
    response = await client.post(
        f"/api/v1/bulk-import/{kind}/preview",
        files={"file": upload(rows, extension)},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


async def confirm(client, headers, batch, kind="leads"):
    response = await client.post(
        f"/api/v1/bulk-import/{kind}/confirm", json={"batch_id": batch["batch_id"]}, headers=headers
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
async def test_general_import_normal_creation_and_idempotency(
    client, mock_db, owner_headers, extension
):
    master = await _lead_master_data(mock_db)
    row = _lead_payload(master, comment="Call tomorrow", next_follow_up_date="30-09-2026")
    batch = await preview(client, owner_headers, [row], extension=extension)
    assert batch["valid"] == 1
    assert await mock_db.leads.count_documents({}) == 0
    result = await confirm(client, owner_headers, batch)
    assert result["imported"] == 1
    lead = await mock_db.leads.find_one({})
    assert lead["stage"] == "fresh"
    assert lead["lead_code"].startswith("AFS-")
    assert await mock_db.lead_notes.count_documents({}) == 1
    assert await mock_db.lead_activities.count_documents({}) >= 1
    assert (await confirm(client, owner_headers, batch))["imported"] == 1
    assert (await preview(client, owner_headers, [row], extension=extension))[
        "state"
    ] == "completed"
    assert await mock_db.leads.count_documents({}) == 1


async def test_partial_invalid_and_duplicate_rows(client, mock_db, owner_headers):
    master = await _lead_master_data(mock_db)
    valid = _lead_payload(master)
    rows = [
        valid,
        valid,
        _lead_payload(master, mobile="123"),
        _lead_payload(master, mobile="9611111112", email="bad"),
        _lead_payload(master, mobile="9611111113", next_follow_up_date="31-02-2026"),
        _lead_payload(master, mobile="9611111114", full_name=""),
    ]
    batch = await preview(client, owner_headers, rows)
    assert (batch["total"], batch["valid"], batch["duplicate"], batch["invalid"]) == (6, 1, 1, 4)
    assert all(row["error"] for row in batch["rows"][1:])
    result = await confirm(client, owner_headers, batch)
    assert (result["imported"], result["duplicate"], result["invalid"]) == (1, 1, 4)
    # A different file still uses the company's active-mobile duplicate rule.
    other = await preview(client, owner_headers, [{**valid, "remarks": "Another file"}])
    assert other["duplicate"] == 1
    assert "AFS-" not in other["rows"][0]["error"]


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
async def test_insurance_import_uses_workflow_and_existing_customer_rules(
    client, mock_db, owner_headers, extension
):
    await _seed_workflow_definitions(mock_db)
    product = await _insurance_product_with_schema(mock_db, product_name="Health", docs=["PAN"])
    row = _payload(product, alternate_mobile="9876543211", nominee_dob="2000-01-01")
    row.pop("stage")
    batch = await preview(client, owner_headers, [row, row], "insurance", extension)
    assert batch["valid"] == 2  # Manual Insurance allows multiple applications per customer.
    assert await mock_db.application_workflows.count_documents({}) == 0
    result = await confirm(client, owner_headers, batch, "insurance")
    assert (result["imported"], result["duplicate"]) == (2, 0)
    assert (
        await mock_db.application_workflows.count_documents(
            {"case_type": "insurance", "current_status": "fresh_lead"}
        )
        == 2
    )
    assert await mock_db.customers.count_documents({}) == 1
    assert await mock_db.leads.count_documents({}) == 0
    assert await mock_db.application_notes.count_documents({}) == 2
    assert (await confirm(client, owner_headers, batch, "insurance"))["imported"] == 2
    assert await mock_db.applications.count_documents({}) == 2


async def test_insurance_invalid_gender_mobile_and_category(client, mock_db, owner_headers):
    product = await _insurance_product_with_schema(mock_db, product_name="Health", docs=["PAN"])
    row = _payload(product)
    row.pop("stage")
    rows = [
        {**row, "gender": "unknown"},
        {**row, "mobile": "123"},
        {**row, "insurance_category_id": "missing"},
        {**row, "email": "bad"},
    ]
    batch = await preview(client, owner_headers, rows, "insurance")
    assert batch["invalid"] == 4
    assert (await confirm(client, owner_headers, batch, "insurance"))["imported"] == 0


@pytest.mark.parametrize(
    "filename,content,mime",
    [
        ("test.xls", b"old workbook", "application/vnd.ms-excel"),
        ("test.exe", b"code", "application/octet-stream"),
        (
            "test.xlsx",
            b"not a workbook",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        ("test.csv", b"Full Name\nPerson", "text/csv"),
        ("test.csv", b"\xff\xfe", "text/csv"),
        ("test.csv", b"\x00", "text/csv"),
        ("test.csv", b"x" * (5 * 1024 * 1024 + 1), "text/csv"),
        ("test.csv", b"Full Name", "image/png"),
    ],
    ids=[
        "xls",
        "executable",
        "malformed-xlsx",
        "missing-columns",
        "encoding",
        "binary",
        "oversized",
        "mime",
    ],
)
async def test_reject_unsafe_or_malformed_files(client, owner_headers, filename, content, mime):
    response = await client.post(
        "/api/v1/bulk-import/leads/preview",
        files={"file": (filename, content, mime)},
        headers=owner_headers,
    )
    assert response.status_code == 422, response.text
    assert "Traceback" not in response.text


@pytest.mark.parametrize("kind", ["leads", "insurance"])
async def test_import_endpoints_deny_unauthorized_employee(client, employee_headers, kind):
    assert (
        await client.get(f"/api/v1/bulk-import/{kind}/sample", headers=employee_headers)
    ).status_code == 403
    assert (
        await client.post(
            f"/api/v1/bulk-import/{kind}/preview",
            files={"file": ("file.csv", b"x", "text/csv")},
            headers=employee_headers,
        )
    ).status_code == 403
    assert (
        await client.post(
            f"/api/v1/bulk-import/{kind}/confirm",
            json={"batch_id": "a" * 64},
            headers=employee_headers,
        )
    ).status_code == 403


@pytest.mark.parametrize("kind", ["leads", "insurance"])
async def test_sample_matches_import_and_schema(client, mock_db, owner_headers, kind):
    if kind == "leads":
        await _lead_master_data(mock_db)
    else:
        await _insurance_product_with_schema(mock_db, product_name="Health", docs=["PAN"])
    response = await client.get(f"/api/v1/bulk-import/{kind}/sample", headers=owner_headers)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    content = base64.b64decode(data["content"])
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Leads", "Instructions", "Lookups"]
    response = await client.post(
        f"/api/v1/bulk-import/{kind}/preview",
        files={
            "file": (
                data["filename"],
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=owner_headers,
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["valid"] == 1


async def test_formula_and_row_limit(client, mock_db, owner_headers):
    master = await _lead_master_data(mock_db)
    row = _lead_payload(master, remarks="=WEBSERVICE(1)")
    batch = await preview(client, owner_headers, [row], extension="xlsx")
    assert batch["invalid"] == 1
    response = await client.post(
        "/api/v1/bulk-import/leads/preview",
        files={"file": upload([row] * 501)},
        headers=owner_headers,
    )
    assert response.status_code == 422


async def test_revalidation_and_partial_creation_failure(
    client, mock_db, owner_headers, monkeypatch
):
    master = await _lead_master_data(mock_db)
    batch = await preview(
        client, owner_headers, [_lead_payload(master), _lead_payload(master, mobile="9611111112")]
    )
    original = LeadService.create_lead

    async def fail_second(self, payload, actor):
        if payload.mobile.endswith("2"):
            raise RuntimeError("Sensitive internal content")
        return await original(self, payload, actor)

    monkeypatch.setattr(LeadService, "create_lead", fail_second)
    result = await confirm(client, owner_headers, batch)
    assert (result["imported"], result["failed"]) == (1, 1)
    assert "Sensitive" not in str(result)
    assert (await confirm(client, owner_headers, batch))["failed"] == 1
    assert await mock_db.leads.count_documents({}) == 1


async def test_batch_binding_expiry_and_interruption(client, mock_db, mock_redis, owner_headers):
    master = await _lead_master_data(mock_db)
    batch = await preview(client, owner_headers, [_lead_payload(master)])
    response = await client.post(
        "/api/v1/bulk-import/insurance/confirm",
        json={"batch_id": batch["batch_id"]},
        headers=owner_headers,
    )
    assert response.status_code == 404
    await mock_redis.set(f"bulk-import:{batch['batch_id']}:lock", "1", ex=100)
    response = await client.post(
        "/api/v1/bulk-import/leads/confirm",
        json={"batch_id": batch["batch_id"]},
        headers=owner_headers,
    )
    assert response.status_code == 409
    assert await mock_db.leads.count_documents({}) == 0
    await mock_redis.delete(f"bulk-import:{batch['batch_id']}")
    response = await client.post(
        "/api/v1/bulk-import/leads/confirm",
        json={"batch_id": batch["batch_id"]},
        headers=owner_headers,
    )
    assert response.status_code == 404
