import csv
import io
from datetime import timedelta

import pytest
from app.features.auth.models import User
from app.features.daily_operations.catalog import install_catalog
from app.features.daily_operations.service import DailyOperationsService
from app.features.daily_operations.transfers import COLUMNS, spreadsheet
from app.utils.datetime import utc_now
from bson import ObjectId
from openpyxl import Workbook, load_workbook
from test_leads import _create_employee, _create_lead, _lead_master_data
from test_workflow import _login

BASE = "/api/v1/daily-operations"


@pytest.fixture
async def follow_payload(client, mock_db, owner_headers):
    await install_catalog(mock_db)
    master = await _lead_master_data(mock_db)
    lead = await _create_lead(client, owner_headers, master)
    owner = await mock_db.users.find_one({"role": "owner"})
    await mock_db.leads.update_one(
        {"_id": ObjectId(lead["id"])}, {"$set": {"assigned_to": str(owner["_id"])}}
    )
    return {
        "date": "2026-10-25",
        "lead_id": lead["id"],
        "bank": "Test Bank",
        "status": "NEXT FOLLOW UP",
        "remarks": "Initial attempt",
        "remarks_date": "2026-10-25",
        "next_action": "Review documents",
        "due_date": "2026-10-26",
    }


async def create(client, headers, area, payload):
    result = await client.post(f"{BASE}/{area}", json=payload, headers=headers)
    assert result.status_code == 200, result.text
    return result.json()["data"]


async def test_follow_up_history_concurrency_and_filters(
    client, mock_db, owner_headers, follow_payload
):
    record = await create(client, owner_headers, "follow-ups", follow_payload)
    changed = {**follow_payload, "remarks": "Second attempt", "version": 1}
    response = await client.patch(
        f"{BASE}/follow-ups/{record['id']}", json=changed, headers=owner_headers
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["version"] == 2
    assert (
        await client.patch(
            f"{BASE}/follow-ups/{record['id']}", json=changed, headers=owner_headers
        )
    ).status_code == 409
    detail = (
        await client.get(f"{BASE}/follow-ups/{record['id']}", headers=owner_headers)
    ).json()["data"]
    assert detail["history"][1]["before"]["remarks"] == "Initial attempt"
    assert len(detail["history"]) == 2
    for suffix, expected in [
        ("?bank=Test+Bank", 1),
        ("?search=documents", 1),
        ("?date_from=2027-01-01", 0),
        ("?sort=status&page_size=1", 1),
    ]:
        response = await client.get(f"{BASE}/follow-ups{suffix}", headers=owner_headers)
        assert response.status_code == 200, response.text
        assert response.json()["data"]["total"] == expected
    assert (
        await client.get(f"{BASE}/follow-ups?sort=unknown", headers=owner_headers)
    ).status_code == 422


@pytest.mark.parametrize(
    "changes",
    [
        {"next_action": ""},
        {"due_date": None},
        {"status": "REJECTED WITH REASON", "due_date": None},
        {"status": "DISBURSE", "due_date": None},
        {"status": "UNKNOWN"},
        {"rejection_reason": "Contradictory reason"},
        {"assigned_to": "forged employee"},
    ],
)
async def test_server_validation(client, owner_headers, follow_payload, changes):
    assert (
        await client.post(
            f"{BASE}/follow-ups",
            headers=owner_headers,
            json={**follow_payload, **changes},
        )
    ).status_code == 422


async def test_rejected_transition_preserves_old_due_date(
    client, owner_headers, follow_payload
):
    record = await create(client, owner_headers, "follow-ups", follow_payload)
    response = await client.patch(
        f"{BASE}/follow-ups/{record['id']}",
        headers=owner_headers,
        json={
            **follow_payload,
            "version": 1,
            "status": "REJECTED WITH REASON",
            "due_date": None,
            "rejection_reason": "Documents not supplied",
        },
    )
    assert response.status_code == 200
    assert response.json()["data"]["history"][-1]["before"]["due_date"] == "2026-10-26"


async def grant(mock_db, employee, area, actions):
    permission = await mock_db.permissions.find_one(
        {"module": "daily_operations", "resource": area}
    )
    role_id = ObjectId()
    await mock_db.employee_roles.insert_one(
        {
            "employee_id": employee["id"],
            "role_id": str(role_id),
            "is_deleted": False,
            "status": "active",
        }
    )
    await mock_db.role_permissions.insert_one(
        {
            "role_id": str(role_id),
            "permission_id": str(permission["_id"]),
            "granted_actions": actions,
            "denied_actions": [],
            "department_ids": None,
            "branch_ids": None,
            "is_deleted": False,
            "status": "active",
        }
    )


async def test_permission_denial_and_issue_scope(
    client, mock_db, owner_headers, master_data
):
    await install_catalog(mock_db)
    employee = await _create_employee(client, owner_headers, master_data)
    headers = await _login(client, "9511111111", "InitialPass1!")
    assert (await client.get(f"{BASE}/issues", headers=headers)).status_code == 200
    assert (await client.get(f"{BASE}/issues")).status_code == 401
    await grant(mock_db, employee, "issues", ["view", "create", "edit"])
    payload = {"date": "2026-10-25", "inquiry": "Test operational question"}
    own = await create(client, headers, "issues", payload)
    other = await create(client, owner_headers, "issues", payload)
    response = await client.get(f"{BASE}/issues", headers=headers)
    assert response.status_code == 200, response.text
    assert [r["id"] for r in response.json()["data"]["items"]] == [own["id"]]
    assert (
        await client.get(f"{BASE}/issues/{other['id']}", headers=headers)
    ).status_code == 404
    response = await client.patch(
        f"{BASE}/issues/{own['id']}",
        json={
            **payload,
            "status": "RESOLVED",
            "resolution_notes": "Resolved",
            "version": 1,
        },
        headers=headers,
    )
    assert response.status_code == 200
    response = await client.patch(
        f"{BASE}/issues/{own['id']}",
        json={
            **payload,
            "status": "REOPENED",
            "assigned_to": str(ObjectId()),
            "version": 2,
        },
        headers=headers,
    )
    assert response.status_code == 403


async def test_issue_resolution_reopen_retains_history(client, owner_headers, mock_db):
    await install_catalog(mock_db)
    payload = {
        "date": "2026-10-25",
        "inquiry": "Internal test issue",
        "priority": "high",
    }
    record = await create(client, owner_headers, "issues", payload)
    for version, status, notes in [
        (1, "IN PROGRESS", ""),
        (2, "RESOLVED", "Resolution recorded"),
        (3, "REOPENED", ""),
    ]:
        response = await client.patch(
            f"{BASE}/issues/{record['id']}",
            headers=owner_headers,
            json={
                **payload,
                "version": version,
                "status": status,
                "resolution_notes": notes,
            },
        )
        assert response.status_code == 200, response.text
    result = response.json()["data"]
    assert result["resolved_at"] is None
    assert result["history"][-1]["before"]["resolution_notes"] == "Resolution recorded"


async def preview(client, headers, rows, kind="follow-ups", format="csv"):
    columns = COLUMNS[kind]
    content = spreadsheet(
        list(columns.values()), [[row.get(k) for k in columns] for row in rows], format
    )
    if format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(columns.values())
        writer.writerows([[row.get(k) for k in columns] for row in rows])
        content = stream.getvalue().encode("utf-8-sig")
    response = await client.post(
        f"{BASE}/transfers/{kind}/preview",
        headers=headers,
        files={"file": (f"test.{format}", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


@pytest.mark.parametrize("format", ["csv", "xlsx"])
async def test_import_preview_confirm_retry_and_export(
    client, mock_db, owner_headers, follow_payload, format
):
    batch = await preview(
        client,
        owner_headers,
        [follow_payload, {**follow_payload, "next_action": ""}],
        format=format,
    )
    assert len(batch["rows"][1]["errors"]) == 1
    assert await mock_db.daily_follow_ups.count_documents({}) == 0
    url = f"{BASE}/transfers/{batch['id']}/confirm"
    assert (
        await client.post(url, json={"confirmed": False}, headers=owner_headers)
    ).status_code == 422
    response = await client.post(url, json={"confirmed": True}, headers=owner_headers)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["inserted"] == 1
    assert response.json()["data"]["rejected"] == 1
    retry = await client.post(url, json={"confirmed": True}, headers=owner_headers)
    assert retry.json()["data"]["skipped"] == 1
    response = await client.get(
        f"{BASE}/follow-ups/export?format={format}", headers=owner_headers
    )
    assert response.status_code == 200, response.text
    if format == "xlsx":
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.active["A1"].value == "DATE"
        workbook.close()
    else:
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert rows[0][0] == "DATE" and len(rows) == 2
    assert (
        await mock_db.audit_logs.count_documents({"event_type": "daily_import_commit"})
        == 2
    )


async def test_import_revalidates_relationships_and_batch_ownership(
    client, mock_db, owner_headers, follow_payload, master_data
):
    batch = await preview(client, owner_headers, [follow_payload])
    await _create_employee(client, owner_headers, master_data)
    headers = await _login(client, "9511111111", "InitialPass1!")
    response = await client.post(
        f"{BASE}/transfers/{batch['id']}/confirm",
        json={"confirmed": True},
        headers=headers,
    )
    assert response.status_code == 404
    await mock_db.leads.update_one(
        {"_id": ObjectId(follow_payload["lead_id"])}, {"$set": {"is_deleted": True}}
    )
    response = await client.post(
        f"{BASE}/transfers/{batch['id']}/confirm",
        json={"confirmed": True},
        headers=owner_headers,
    )
    assert response.json()["data"]["rejected"] == 1
    assert await mock_db.daily_follow_ups.count_documents({}) == 0


@pytest.mark.parametrize(
    "filename,content",
    [
        ("bad.xlsx", b"not a zip"),
        ("macro.xlsm", b"bad"),
        ("test.csv", b"bad,headings\n1,2"),
        ("test.csv", b"\x00binary"),
    ],
)
async def test_invalid_files(client, owner_headers, filename, content):
    response = await client.post(
        f"{BASE}/transfers/issues/preview",
        headers=owner_headers,
        files={"file": (filename, content)},
    )
    assert response.status_code == 422, response.text


async def test_sparse_workbook_and_formula_injection(
    client, owner_headers, follow_payload
):
    workbook = Workbook()
    workbook.active["A1000000"] = "sparse"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    response = await client.post(
        f"{BASE}/transfers/follow-ups/preview",
        headers=owner_headers,
        files={"file": ("sparse.xlsx", stream.getvalue())},
    )
    assert response.status_code == 422
    batch = await preview(
        client, owner_headers, [{**follow_payload, "bank": "=SUM(1,2)"}]
    )
    assert batch["rows"][0]["errors"]
    await create(
        client, owner_headers, "follow-ups", {**follow_payload, "bank": "=untrusted"}
    )
    exported = await client.get(f"{BASE}/follow-ups/export", headers=owner_headers)
    assert "'=untrusted" in exported.text


async def test_sessions_group_rotations_and_preserve_original_correction(
    client, mock_db, owner_headers, master_data
):
    await install_catalog(mock_db)
    employee = await _create_employee(client, owner_headers, master_data)
    user_id = employee["user_id"]
    now = utc_now()
    family = "isolated-test-family"
    for status in ["rotated", "logged_out"]:
        await mock_db.sessions.insert_one(
            {
                "token_id": str(ObjectId()),
                "family_id": family,
                "user_id": user_id,
                "login_at": now - timedelta(hours=2),
                "logout_at": now - timedelta(hours=1)
                if status == "logged_out"
                else None,
                "updated_at": now,
                "status": status,
                "is_deleted": False,
            }
        )
    response = await client.get(f"{BASE}/sessions", headers=owner_headers)
    assert response.status_code == 200, response.text
    row = next(
        item for item in response.json()["data"]["items"] if item["id"] == family
    )
    assert row["recorded_seconds"] == 3600
    correction = {
        "login_at": (now - timedelta(hours=3)).isoformat(),
        "logout_at": (now - timedelta(hours=1)).isoformat(),
        "reason": "Test administrative correction",
        "version": 0,
    }
    response = await client.post(
        f"{BASE}/sessions/{family}/correction", headers=owner_headers, json=correction
    )
    assert response.status_code == 200, response.text
    assert (
        await client.post(
            f"{BASE}/sessions/{family}/correction",
            headers=owner_headers,
            json=correction,
        )
    ).status_code == 409
    response = await client.get(f"{BASE}/sessions", headers=owner_headers)
    row = next(
        item for item in response.json()["data"]["items"] if item["id"] == family
    )
    assert row["corrected"] and row["recorded_seconds"] == 7200
    assert (
        await mock_db.sessions.count_documents(
            {"family_id": family, "login_at": now - timedelta(hours=2)}
        )
        == 2
    )
    correction["login_at"] = now.isoformat()
    assert (
        await client.post(
            f"{BASE}/sessions/{family}/correction",
            headers=owner_headers,
            json=correction,
        )
    ).status_code == 422


async def test_permissions_do_not_grant_session_corrections_to_employee(
    client, mock_db, owner_headers, master_data
):
    await install_catalog(mock_db)
    employee = await _create_employee(client, owner_headers, master_data)
    await grant(mock_db, employee, "sessions", ["view", "edit"])
    headers = await _login(client, "9511111111", "InitialPass1!")
    now = utc_now()
    response = await client.post(
        f"{BASE}/sessions/fake/correction",
        headers=headers,
        json={
            "login_at": (now - timedelta(hours=2)).isoformat(),
            "logout_at": (now - timedelta(hours=1)).isoformat(),
            "reason": "Unauthorized correction",
        },
    )
    assert response.status_code == 403


async def test_catalog_is_idempotent_and_does_not_grant_access(mock_db):
    await install_catalog(mock_db)
    await install_catalog(mock_db)
    assert (
        await mock_db.permissions.count_documents({"module": "daily_operations"}) == 4
    )
    assert await mock_db.role_permissions.count_documents({}) == 0
    assert await mock_db.nav_items.count_documents({"key": {"$regex": "^daily_"}}) == 4
    outsider = User(_id=str(ObjectId()), mobile="9000000000", role="customer")
    from app.core.exceptions import ForbiddenError

    with pytest.raises(ForbiddenError):
        await DailyOperationsService(mock_db).authorize(outsider, "issues")


async def seed_case(mock_db, lead_id, *, disbursed=False):
    from app.features.customer.models import Application
    from app.features.workflow_engine.models import ApplicationWorkflow, LoanCaseDetails

    lead = await mock_db.leads.find_one({"_id": ObjectId(lead_id)})
    customer_id = str(ObjectId())
    application = Application(
        application_code="TEST-APP",
        user_id=str(ObjectId()),
        product_category="loan",
        product_id=lead["product_id"],
        form_definition_id=str(ObjectId()),
        lead_id=lead_id,
        customer_id=customer_id,
    )
    application_id = str(
        (
            await mock_db.applications.insert_one(
                application.model_dump(exclude={"id"})
            )
        ).inserted_id
    )
    details = LoanCaseDetails(
        bank_nbfc_name="Test Bank",
        loan_type="Test Loan",
        offered_amount=1000,
        disbursed_amount=900 if disbursed else None,
        disbursed_at=utc_now() if disbursed else None,
        disbursed_reference="TEST-REF" if disbursed else None,
    )
    case = ApplicationWorkflow(
        case_code="TEST-CASE",
        case_type="loan",
        application_id=application_id,
        customer_id=customer_id,
        product_id=lead["product_id"],
        product_category="loan",
        current_status="disbursed" if disbursed else "credit_evaluation",
        assigned_to=lead["assigned_to"],
        moved_to_loan_management_at=utc_now(),
        loan_details=details,
    )
    return str(
        (
            await mock_db.application_workflows.insert_one(
                case.model_dump(exclude={"id"})
            )
        ).inserted_id
    )


async def test_verified_disbursement_and_relationship_integrity(
    client, mock_db, owner_headers, follow_payload
):
    case_id = await seed_case(mock_db, follow_payload["lead_id"])
    payload = {
        **follow_payload,
        "case_id": case_id,
        "status": "DISBURSE",
        "due_date": None,
    }
    assert (
        await client.post(f"{BASE}/follow-ups", headers=owner_headers, json=payload)
    ).status_code == 422
    await mock_db.application_workflows.update_one(
        {"_id": ObjectId(case_id)},
        {
            "$set": {
                "current_status": "disbursed",
                "loan_details.disbursed_amount": 900,
                "loan_details.disbursed_at": utc_now(),
                "loan_details.disbursed_reference": "TEST-REF",
            }
        },
    )
    await create(client, owner_headers, "follow-ups", payload)
    await mock_db.applications.update_many({}, {"$set": {"lead_id": str(ObjectId())}})
    assert (
        await client.post(f"{BASE}/follow-ups", headers=owner_headers, json=payload)
    ).status_code == 422


async def test_progress_uses_actual_amounts_and_imports_only_follow_up_actions(
    client, mock_db, owner_headers, follow_payload
):
    case_id = await seed_case(mock_db, follow_payload["lead_id"])
    response = await client.get(f"{BASE}/progress", headers=owner_headers)
    assert response.status_code == 200, response.text
    item = response.json()["data"]["items"][0]
    assert item["approved_amount"] == 1000 and item["disbursed_amount"] is None
    batch = await preview(
        client,
        owner_headers,
        [
            {
                "case_id": case_id,
                "date": "2026-10-25",
                "next_action": "Contact bank",
                "due_date": "2026-10-26",
            }
        ],
        kind="progress",
    )
    result = await client.post(
        f"{BASE}/transfers/{batch['id']}/confirm",
        headers=owner_headers,
        json={"confirmed": True},
    )
    assert result.json()["data"]["inserted"] == 1, result.text
    response = await client.get(f"{BASE}/progress", headers=owner_headers)
    item = response.json()["data"]["items"][0]
    assert item["next_action"] == "Contact bank" and item["due_date"] == "2026-10-26"
    assert (await mock_db.application_workflows.find_one({"_id": ObjectId(case_id)}))[
        "current_status"
    ] == "credit_evaluation"
    assert (
        await client.get(f"{BASE}/progress?old_days=10", headers=owner_headers)
    ).json()["data"]["total"] == 0


async def test_exact_duplicate_rows_skip_and_row_limit_rejects(
    client, owner_headers, follow_payload
):
    batch = await preview(client, owner_headers, [follow_payload, follow_payload])
    assert batch["rows"][1]["duplicate"]
    response = await client.post(
        f"{BASE}/transfers/{batch['id']}/confirm",
        headers=owner_headers,
        json={"confirmed": True},
    )
    assert (
        response.json()["data"]["inserted"] == 1
        and response.json()["data"]["skipped"] == 1
    )
    too_many = spreadsheet(
        list(COLUMNS["follow-ups"].values()),
        [[follow_payload.get(k) for k in COLUMNS["follow-ups"]]] * 501,
        "csv",
    )
    assert (
        await client.post(
            f"{BASE}/transfers/follow-ups/preview",
            headers=owner_headers,
            files={"file": ("large.csv", too_many)},
        )
    ).status_code == 422


async def test_mapping_and_envelope_limits(client, owner_headers):
    response = await client.post(
        f"{BASE}/transfers/issues/preview",
        headers=owner_headers,
        files={
            "file": (
                "mapped.csv",
                b"My Date,My Question\r\n2026-10-25,Internal test\r\n",
            )
        },
        data={"mapping": '{"My Date":"date","My Question":"inquiry"}'},
    )
    assert (
        response.status_code == 200 and not response.json()["data"]["rows"][0]["errors"]
    ), response.text
    response = await client.post(
        f"{BASE}/transfers/issues/preview",
        headers={**owner_headers, "Content-Length": "99999999"},
        content=b"test",
    )
    assert response.status_code == 413


async def test_import_issue_and_session_correction(
    client, mock_db, owner_headers, master_data
):
    await install_catalog(mock_db)
    issue = await preview(
        client,
        owner_headers,
        [{"date": "2026-10-25", "inquiry": "Test imported issue"}],
        kind="issues",
        format="xlsx",
    )
    response = await client.post(
        f"{BASE}/transfers/{issue['id']}/confirm",
        headers=owner_headers,
        json={"confirmed": True},
    )
    assert response.json()["data"]["inserted"] == 1, response.text
    employee = await _create_employee(client, owner_headers, master_data)
    family = "correction-import-test"
    now = utc_now()
    await mock_db.sessions.insert_one(
        {
            "token_id": str(ObjectId()),
            "family_id": family,
            "user_id": employee["user_id"],
            "login_at": now - timedelta(hours=2),
            "logout_at": None,
            "updated_at": now,
            "status": "active",
            "is_deleted": False,
        }
    )
    batch = await preview(
        client,
        owner_headers,
        [
            {
                "family_id": family,
                "login_at": (now - timedelta(hours=2)).isoformat(),
                "logout_at": (now - timedelta(hours=1)).isoformat(),
                "reason": "Confirmed missing logout",
                "version": 0,
            }
        ],
        kind="sessions",
    )
    url = f"{BASE}/transfers/{batch['id']}/confirm"
    assert (
        await client.post(url, headers=owner_headers, json={"confirmed": True})
    ).json()["data"]["updated"] == 1
    assert (
        await client.post(url, headers=owner_headers, json={"confirmed": True})
    ).json()["data"]["skipped"] == 1


async def test_follow_up_assignment_follows_lead_without_leaking_to_previous_owner(
    client, mock_db, owner_headers, follow_payload, master_data
):
    from app.features.access_control.models import Permission

    employee = await _create_employee(client, owner_headers, master_data)
    await grant(mock_db, employee, "follow-ups", ["view"])
    permission = Permission(module="leads", resource="leads", actions=["view"])
    permission_id = (
        await mock_db.permissions.insert_one(permission.model_dump(exclude={"id"}))
    ).inserted_id
    role = await mock_db.employee_roles.find_one({"employee_id": employee["id"]})
    await mock_db.role_permissions.insert_one(
        {
            "role_id": role["role_id"],
            "permission_id": str(permission_id),
            "granted_actions": ["view"],
            "is_deleted": False,
        }
    )
    record = await create(client, owner_headers, "follow-ups", follow_payload)
    headers = await _login(client, "9511111111", "InitialPass1!")
    assert (await client.get(f"{BASE}/follow-ups", headers=headers)).json()["data"][
        "total"
    ] == 0
    await mock_db.leads.update_one(
        {"_id": ObjectId(follow_payload["lead_id"])},
        {"$set": {"assigned_to": employee["id"]}},
    )
    assert (await client.get(f"{BASE}/follow-ups", headers=headers)).json()["data"][
        "total"
    ] == 1
    assert (
        await client.get(f"{BASE}/follow-ups/{record['id']}", headers=headers)
    ).status_code == 200
    await mock_db.leads.update_one(
        {"_id": ObjectId(follow_payload["lead_id"])},
        {"$set": {"assigned_to": str(ObjectId())}},
    )
    assert (
        await client.get(f"{BASE}/follow-ups/{record['id']}", headers=headers)
    ).status_code == 403


async def test_import_updates_require_opt_in_and_version_and_retry_safely(
    client, mock_db, owner_headers, follow_payload
):
    record = await create(client, owner_headers, "follow-ups", follow_payload)
    row = {
        **follow_payload,
        "record_id": record["id"],
        "version": 1,
        "remarks": "Imported update",
    }
    disabled = await preview(client, owner_headers, [row])
    assert "disabled" in disabled["rows"][0]["errors"][0]["message"]
    columns = COLUMNS["follow-ups"]
    content = spreadsheet(
        list(columns.values()), [[row.get(k) for k in columns]], "csv"
    )
    response = await client.post(
        f"{BASE}/transfers/follow-ups/preview",
        headers=owner_headers,
        files={"file": ("update.csv", content)},
        data={"allow_updates": "true"},
    )
    assert response.status_code == 200, response.text
    batch = response.json()["data"]
    assert not batch["rows"][0]["errors"], batch
    url = f"{BASE}/transfers/{batch['id']}/confirm"
    assert (
        await client.post(url, headers=owner_headers, json={"confirmed": True})
    ).json()["data"]["updated"] == 1
    assert (
        await client.post(url, headers=owner_headers, json={"confirmed": True})
    ).json()["data"]["skipped"] == 1
    stored = await mock_db.daily_follow_ups.find_one({"_id": ObjectId(record["id"])})
    assert stored["remarks"] == "Imported update" and len(stored["history"]) == 2
    assert "record_id" not in stored


async def test_internal_notifications_are_authorized_and_retry_safe(
    client, mock_db, owner_headers, follow_payload
):
    from app.features.daily_operations.notifications import deliver_notifications
    from app.utils.datetime import to_ist

    today = to_ist(utc_now()).date().isoformat()
    await create(
        client,
        owner_headers,
        "follow-ups",
        {**follow_payload, "date": today, "remarks_date": today, "due_date": today},
    )
    first = await deliver_notifications(mock_db)
    assert first["created"] == 1
    assert (await deliver_notifications(mock_db))["created"] == 0
    await mock_db.daily_follow_ups.update_many(
        {}, {"$unset": {"notified_due_date": ""}}
    )
    assert (await deliver_notifications(mock_db))["created"] == 0
    await mock_db.leads.update_many({}, {"$set": {"assigned_to": str(ObjectId())}})
    assert (await deliver_notifications(mock_db))["skipped"] == 1
    assert (
        await mock_db.notifications.count_documents(
            {"notification_type": "daily_follow_up_due"}
        )
        == 1
    )


async def test_report_financial_totals_ignore_follow_up_disburse_labels(
    client, mock_db, owner_headers, follow_payload
):
    case_id = await seed_case(mock_db, follow_payload["lead_id"], disbursed=True)
    await create(
        client,
        owner_headers,
        "follow-ups",
        {**follow_payload, "case_id": case_id, "status": "DISBURSE", "due_date": None},
    )
    await create(
        client,
        owner_headers,
        "follow-ups",
        {**follow_payload, "case_id": case_id, "status": "DISBURSE", "due_date": None},
    )
    result = (await client.get(f"{BASE}/progress", headers=owner_headers)).json()[
        "data"
    ]
    assert result["metrics"]["approved_amount"] == 1000
    assert result["metrics"]["disbursed_amount"] == 900
    assert result["total"] == 1


async def test_templates_preserve_headings_and_read_only_customer_identity(
    client, mock_db, owner_headers, follow_payload
):
    template = await client.get(
        f"{BASE}/transfers/follow-ups/template?format=csv", headers=owner_headers
    )
    assert template.status_code == 200
    for heading in [
        "DATE",
        "LOAN TYPE",
        "CX NAME",
        "CX NUMBER",
        "EMPLOYEE NAME",
        "REMARKS DATE",
        "FOLLOW PENDING",
    ]:
        assert heading in template.text
    batch = await preview(
        client, owner_headers, [{**follow_payload, "customer_name": "Forged identity"}]
    )
    assert "read-only" in batch["rows"][0]["errors"][0]["message"]
    record = await create(client, owner_headers, "follow-ups", follow_payload)
    stored = await mock_db.daily_follow_ups.find_one({"_id": ObjectId(record["id"])})
    from app.utils.datetime import to_ist

    assert to_ist(stored["date"]).date().isoformat() == follow_payload["date"]
    assert stored["date"].tzinfo is not None


async def test_concurrent_confirmation_does_not_duplicate_rows(
    client, mock_db, owner_headers, follow_payload
):
    import asyncio

    batch = await preview(client, owner_headers, [follow_payload])
    url = f"{BASE}/transfers/{batch['id']}/confirm"
    responses = await asyncio.gather(
        *[
            client.post(url, headers=owner_headers, json={"confirmed": True})
            for _ in range(2)
        ]
    )
    assert all(response.status_code == 200 for response in responses)
    assert sum(response.json()["data"]["inserted"] for response in responses) == 1
    assert await mock_db.daily_follow_ups.count_documents({}) == 1
